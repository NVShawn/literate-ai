#!/usr/bin/env python3
"""Build the Literate AI conference evolution deck with editable native shapes."""

from __future__ import annotations

import hashlib
import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
OUT = HERE / "literate-ai-evolution-conference.pptx"
VALIDATION = HERE / "validation.json"
ASSETS = HERE / "assets"
W, H, TOTAL = 13.333, 7.5, 14

C = {
    "ink": "101317", "panel": "23282F", "fog": "EEF1F3", "white": "FFFFFF",
    "orange": "FF6B35", "orange2": "FF9B66", "blue": "72B7D6",
    "green": "76B900", "green2": "7BC6A4", "red": "F47C7C",
    "line": "D8DEE3", "steel": "65707C", "muted": "AAB3BC", "night": "07111F",
}
FONT = "Helvetica Neue"
MONO = "Courier New"


def rgb(name: str) -> RGBColor:
    return RGBColor.from_string(C.get(name, name))


def rect(slide, x, y, w, h, fill, radius=False, line=None):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = rgb(fill)
    if line:
        shape.line.color.rgb = rgb(line); shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    return shape


def txt(slide, value, x, y, w, h, size=20, color="ink", bold=False,
        align="left", font=FONT, fit=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]; p.text = value
    p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
    for run in p.runs:
        run.font.name = font; run.font.size = Pt(size); run.font.bold = bold
        run.font.color.rgb = rgb(color)
    if fit:
        tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    return box


def line(slide, x1, y1, x2, y2, color="line", width=2):
    item = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    item.line.color.rgb = rgb(color); item.line.width = Pt(width)
    return item


def pill(slide, label, x, y, w, fill="orange", color="ink"):
    sp = rect(slide, x, y, w, .34, fill, True)
    tf = sp.text_frame; tf.clear(); tf.margin_left = tf.margin_right = Inches(.08)
    tf.margin_top = tf.margin_bottom = 0; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.text = label; p.alignment = PP_ALIGN.CENTER
    for run in p.runs:
        run.font.name = FONT; run.font.size = Pt(10); run.font.bold = True; run.font.color.rgb = rgb(color)


def header(slide, n, title, subtitle="", dark=False, tag="CONFERENCE 2026"):
    fg = "white" if dark else "ink"; muted = "muted" if dark else "steel"
    pill(slide, "LITERATE-AI", .72, .30, 1.35, "white" if dark else "ink", "ink" if dark else "white")
    pill(slide, tag, 2.18, .30, 1.45, "orange", "ink")
    txt(slide, f"{n:02d}", 12.1, .34, .5, .25, 10, fg, True, "right")
    txt(slide, title, .72, .82, 11.8, .92, 27, fg, True, fit=True)
    if subtitle:
        txt(slide, subtitle, .74, 1.84, 11.5, .42, 12, muted, fit=True)
    rect(slide, .72, 2.42, .92, .055, "orange")


def footer(slide, n, dark=False, source="source-notes.md"):
    col = "muted" if dark else "steel"
    txt(slide, source, .72, 7.18, 4.2, .16, 8, col)
    rect(slide, 10.9, 7.23, 1.7, .025, "panel" if dark else "line")
    rect(slide, 10.9, 7.23, 1.7 * n / TOTAL, .025, "orange")


def add_notes(slide, lines):
    slide.notes_slide.notes_text_frame.text = "\n".join(lines)


def card(slide, x, y, w, h, heading, body, fill="white", accent="orange", dark=False):
    rect(slide, x, y, w, h, fill, True, "line" if not dark else "panel")
    rect(slide, x, y, .08, h, accent, True)
    txt(slide, heading, x + .24, y + .22, w - .42, .35, 16, "white" if dark else "ink", True, fit=True)
    txt(slide, body, x + .24, y + .69, w - .42, h - .82, 11.5, "muted" if dark else "steel", fit=True)


def image_crop(slide, path, x, y, w, h):
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))


def build():
    prs = Presentation(); prs.slide_width = Inches(W); prs.slide_height = Inches(H)
    blank = prs.slide_layouts[6]
    def slide(bg="white"):
        s = prs.slides.add_slide(blank); rect(s, 0, 0, W, H, bg); return s

    # 1
    s = slide("ink")
    rect(s, 8.3, 0, 5.03, 7.5, "night")
    for i, (x, y, color) in enumerate([(8.9,1.1,"orange"),(10.7,.8,"blue"),(11.82,2.0,"green"),(9.5,3.2,"orange2"),(11.2,4.0,"blue"),(9.1,5.4,"green2"),(11.8,5.7,"orange")]):
        rect(s, x, y, 1.0 + (i % 3) * .22, .48, color, True)
        if i: line(s, x-.65, y+.24, x, y+.24, "steel", 2)
    pill(s, "LITERATE-AI", .76, .55, 1.4, "white", "ink")
    txt(s, "From readable intent\nto executable evidence", .76, 1.5, 7.1, 1.8, 36, "white", True, fit=True)
    txt(s, "How a specification-led harness grew into a component system for governed software creation.", .80, 3.65, 6.5, 1.0, 18, "fog", fit=True)
    rect(s, .8, 5.25, 1.1, .07, "orange")
    txt(s, "CONFERENCE EDITION · 28 OCTOBER 2026", .8, 5.52, 6.0, .3, 13, "orange2", True)
    txt(s, "Editable deck · repository-backed claims", .8, 6.64, 5.0, .25, 10, "muted")
    add_notes(s, ["Open with the distinction between generation and confidence.", "Literate AI makes readable intent durable and requires generated implementation to earn acceptance.", "Sources: PROJECT.md; README.md."])

    # 2
    s = slide("fog"); header(s, 2, "Code generation made code abundant. Confidence stayed scarce.", "The bottleneck moved from typing implementation to proving that an exact candidate deserves trust.")
    labels = [("PROMPT", "plausible intent"), ("SOURCE", "plausible implementation"), ("TESTS", "plausible self-assessment")]
    for i,(a,b) in enumerate(labels):
        x=.8+i*3.25; card(s,x,3.05,2.7,1.65,a,b,"white","red")
        if i<2: txt(s,"→",x+2.78,3.55,.4,.4,28,"steel",True,"center")
    rect(s,10.58,2.82,1.85,2.15,"ink",True)
    txt(s,"?",11.13,3.05,.75,.8,45,"orange",True,"center")
    txt(s,"release\nconfidence",10.8,4.12,1.4,.55,13,"white",True,"center")
    txt(s,"A candidate can be coherent, tested, and still be wrong about the authority it was meant to implement.",.82,5.55,11.5,.55,20,"ink",True,"center",fit=True)
    footer(s,2); add_notes(s,["Do not frame models as the problem. The missing piece is an independently checkable custody and acceptance boundary.","Source: docs/user/models-and-generation.md."])

    # 3
    s = slide("ink"); header(s,3,"Invert the product: intent is durable; implementation is renewable.","Three different things must remain different.",True)
    for x,h,b,c in [(1.0,"INTENT","Readable behavior\nand constraints","orange"),(4.65,"IMPLEMENTATION","Replaceable source\nand native artifacts","panel"),(8.3,"EVIDENCE","Independent checks\nand current receipts","green2")]:
        card(s,x,3.05,3.05,2.0,h,b,c,"orange" if h!="EVIDENCE" else "green",True)
    txt(s,"AUTHORITY",1.7,5.35,1.6,.3,12,"orange2",True,"center")
    txt(s,"CANDIDATE",5.35,5.35,1.6,.3,12,"blue",True,"center")
    txt(s,"ADMISSION",9.0,5.35,1.6,.3,12,"green2",True,"center")
    footer(s,3,True); add_notes(s,["This is the core thesis. A model can help regenerate implementation, but cannot grant itself authority or acceptance.","Source: PROJECT.md goals 2 and 4."])

    # 4
    s = slide("white"); header(s,4,"A Component DAG keeps authority factored—and still makes art.","Behavior composes without flattening every concern into one prompt or one source tree.")
    nodes=[(6.2,2.8,"Component","orange"),(2.0,4.0,"math","blue"),(4.5,4.4,"physics","green2"),(7.0,4.25,"3D","orange2"),(9.5,3.8,"music","blue"),(10.8,5.25,"video","green")]
    root=nodes[0]
    for x,y,l,c in nodes[1:]: line(s,root[0]+.6,root[1]+.34,x+.55,y+.28,"steel",2)
    for x,y,l,c in nodes:
        rect(s,x,y,1.35,.58,c,True); txt(s,l,x+.08,y+.16,1.19,.22,12,"ink",True,"center",fit=True)
    txt(s,"Each node owns a readable boundary.\nEdges express composition, not copy-and-paste.",.9,5.75,7.5,.72,18,"ink",True,fit=True)
    footer(s,4); add_notes(s,["The DAG is both a technical boundary and a creative palette. Components can be small enough to review and rich enough to collaborate.","Sources: docs/architecture/repository-inheritance.md; labs/05-component-cathedral."])

    # 5
    s=slide("fog"); header(s,5,"The lifecycle is a chain of custody, not a single generate button.","Every transition has an owner, an identity, and a reason to refuse.")
    stages=[("RESOLVE","exact DAG"),("GENERATE","bounded context"),("BUILD","native tools"),("VERIFY","independent oracle"),("RECEIPT","current evidence"),("RELEASE","exact commit")]
    for i,(a,b) in enumerate(stages):
        x=.55+i*2.1; fill="orange" if i==0 else ("green2" if i>=4 else "white")
        rect(s,x,3.15,1.75,1.25,fill,True,"line"); txt(s,a,x+.12,3.42,1.51,.25,12,"ink",True,"center"); txt(s,b,x+.12,3.82,1.51,.22,10,"steel" if fill=="white" else "ink",False,"center")
        if i<5: txt(s,"→",x+1.78,3.53,.3,.3,19,"steel",True,"center")
    txt(s,"A changed specification, Flavor, skill, route, dependency, toolchain, or artifact changes what must be proved.",1.1,5.25,11.1,.55,18,"ink",True,"center",fit=True)
    footer(s,5); add_notes(s,["The exact recipe is broader than source. Reuse is allowed only when the relevant identities remain current.","Sources: docs/user/framework-flow.md; docs/user/caches-packages-publication.md."])

    # 6
    s=slide("white"); header(s,6,"Evolution I · make readable authority executable.","The 0.x line built the grammar, the lifecycle, and the operator front door.")
    eras=[("0.1–0.4","Components\n+ readable specs"),("0.5–0.8","locks, receipts\n+ release protocol"),("0.9","evidence-gated\nframework"),("0.10","create + adopt\noperator journey"),("1.0","stable product\nboundary")]
    line(s,1.0,4.0,12.1,4.0,"line",4)
    for i,(v,b) in enumerate(eras):
        x=1.0+i*2.7; rect(s,x,3.72,.34,.34,"orange" if i<4 else "green",True); txt(s,v,x-.35,2.9,1.1,.3,15,"ink",True,"center"); txt(s,b,x-.66,4.35,1.7,.7,11,"steel",False,"center",fit=True)
    footer(s,6); add_notes(s,["This is a capability narrative, not a claim that every present feature existed at each tag.","Sources: repository tags, CHANGELOG.md, historical roadmap records."])

    # 7
    s=slide("ink"); header(s,7,"Evolution II · make evidence compose at production scale.","The 1.x line broadens native custody, observability, scheduling, and release governance.",True)
    eras=[("1.1","native + multi-language","blue"),("1.2","integration + document pairs","orange2"),("1.3","action DAG + HTML explorer","green2"),("1.3.1","separate patch closure","orange"),("1.4 dev","retained native + public paths","green")]
    for i,(v,b,c) in enumerate(eras):
        x=.75+i*2.48; rect(s,x,3.0,2.05,2.15,"panel",True,"steel"); rect(s,x,3.0,2.05,.12,c); txt(s,v,x+.18,3.35,1.7,.35,20,"white",True); txt(s,b,x+.18,4.05,1.7,.72,12,"muted",False,fit=True)
    txt(s,"Published",.85,5.72,6.0,.3,12,"blue",True); txt(s,"Patch program",8.15,5.72,1.6,.3,12,"orange",True,"center"); txt(s,"Development",10.65,5.72,1.6,.3,12,"green2",True,"center")
    footer(s,7,True); add_notes(s,["Be explicit: v1.3.0 is the latest published tag. 1.3.1 is a separate patch program. main carries 1.4.0 development authority.","Sources: tags; src/literate_ai/version.py; docs/roadmap/release-goals.json."])

    # 8
    s=slide("fog"); header(s,8,"Exact identities turn explanation into an executable contract.","The recipe is inspectable before model egress and re-checkable after native execution.")
    left=["Component revision","selected Flavors","pinned skills","workflow + routing","repository sources"]
    for i,v in enumerate(left):
        y=2.82+i*.66; rect(s,.8,y,2.8,.46,"white",True,"line"); txt(s,v,.98,y+.12,2.45,.2,11,"ink",True)
        line(s,3.62,y+.23,5.05,3.98,"steel",1.4)
    rect(s,5.05,3.25,2.45,1.5,"ink",True); txt(s,"EXACT\nRECIPE",5.42,3.62,1.7,.7,21,"white",True,"center")
    line(s,7.5,3.98,8.35,3.98,"orange",3)
    rect(s,8.35,2.95,3.8,2.05,"green2",True); txt(s,"candidate + evidence",8.75,3.35,3.0,.35,18,"ink",True,"center"); txt(s,"source · build · tests\nacceptance · receipt",8.75,4.05,3.0,.55,12,"ink",False,"center")
    footer(s,8); add_notes(s,["The receipt is compact because it binds larger retained evidence by exact identity. It does not replace that evidence.","Sources: docs/architecture/design-traceability.md; docs/user/framework-flow.md."])

    # 9
    s=slide("white"); header(s,9,"Failure is evidence, too—if ownership stays explicit.","A timeout or transport loss after actuation cannot be wished into a clean retry.")
    card(s,.8,3.0,3.55,2.25,"MODEL OUTPUT","Useful proposal\nNo acceptance authority","fog","blue")
    card(s,4.88,3.0,3.55,2.25,"FAILED ATTEMPT","Retain diagnostics\nPreserve ownership uncertainty","fog","red")
    card(s,8.96,3.0,3.55,2.25,"TRUSTED RESOLUTION","Owned terminal evidence\nor admitted fenced lease","fog","green")
    txt(s,"Retry policy must follow side effects—not optimism.",1.1,5.85,11.1,.45,21,"ink",True,"center")
    footer(s,9); add_notes(s,["This is why the action DAG distinguishes immutable production from mutable package or signing operations.","Sources: CHANGELOG.md action ownership entries; ADR 0044."])

    # 10
    s=slide("ink"); header(s,10,"Observability became a product surface, not a debug afterthought.","The HTML explorer renders effective authority and provenance as one portable file.",True)
    rect(s,.75,2.78,7.3,3.5,"panel",True,"steel")
    for i,(label,color) in enumerate([("repositories","orange"),("Components","blue"),("Flavors","green2"),("skills","orange2")]):
        x=1.15+i*1.62; rect(s,x,3.24,1.3,.5,color,True); txt(s,label,x+.06,3.39,1.18,.18,10,"ink",True,"center",fit=True)
    for i in range(6):
        x=1.25+(i%3)*2.1; y=4.35+(i//3)*.82; rect(s,x,y,1.5,.42,"night",True,"steel"); line(s,x+1.5,y+.21,7.2,4.95,"steel",1)
    card(s,8.55,2.78,3.85,1.35,"IT IS","A provenance-bound authority graph\nwith filters, focus, and source excerpts","panel","green2",True)
    card(s,8.55,4.42,3.85,1.35,"IT IS NOT","A general IDE, symbol index,\nor go-to-definition engine","panel","red",True)
    footer(s,10,True); add_notes(s,["The boundary matters: calling it a source browser would overstate the implementation.","Source: docs/user/html-repository-explorer.md."])

    # 11
    s=slide("fog"); header(s,11,"Six labs make increasingly questionable life choices.","No KPI dashboards were harmed. Every lab is offline, free-standing, and inspectable.")
    image_crop(s,ASSETS/"00-grumpy-counter.png",.72,2.82,3.7,2.35)
    image_crop(s,ASSETS/"03-impossible-furniture.png",4.82,2.82,3.7,2.35)
    image_crop(s,ASSETS/"05-component-cathedral.png",8.92,2.82,3.7,2.35)
    for x,a,b in [(.72,"00","grumpy math"),(4.82,"03","impossible 3D"),(8.92,"05","multimedia cathedral")]:
        pill(s,a,x,5.42,.52,"orange","ink"); txt(s,b,x+.65,5.47,2.8,.22,12,"ink",True)
    txt(s,"math → physics → audio/music → 3D → video → composed component art",1.0,6.2,11.3,.35,17,"ink",True,"center")
    footer(s,11); add_notes(s,["The source is DRY, but each generated index.html is free-standing and works without a server or network.","Source: labs/README.md and labs/scripts/verify_labs.py."])

    # 12
    s=slide("ink"); header(s,12,"The cathedral is the argument: rigor can wear better clothes.","Yes, it has a pipe organ. No, finance did not ask for it.",True)
    image_crop(s,ASSETS/"05-component-cathedral.png",.72,2.72,8.25,3.95)
    card(s,9.38,2.72,3.2,1.15,"READABLE","Select a node and inspect the authority it owns.","panel","orange",True)
    card(s,9.38,4.08,3.2,1.15,"REPLACEABLE","Change one capability without rewriting the whole artifact.","panel","blue",True)
    card(s,9.38,5.44,3.2,1.15,"PROVABLE","Capture a bounded video artifact; keep audio opt-in.","panel","green2",True)
    footer(s,12,True); add_notes(s,["This is intentionally cheeky. The point is that component boundaries can support expressive work, not only business workflows.","Source: labs/05-component-cathedral."])

    # 13
    s=slide("white"); header(s,13,"Release posture is part of the feature—not a footnote.","Literate AI distinguishes published evidence, patch closure, current development, and open qualification.")
    bands=[("PUBLISHED","1.3.0 exact tag","green2"),("PATCH","1.3.1 explicit goal ledger","orange"),("DEVELOPMENT","1.4.0 on main","blue"),("OPEN","FreeBSD + complete retained/public paths","red")]
    for i,(a,b,c) in enumerate(bands):
        y=2.75+i*.85; rect(s,.9,y,2.0,.56,c,True); txt(s,a,1.08,y+.17,1.64,.2,11,"ink",True,"center"); txt(s,b,3.25,y+.13,8.8,.28,16,"ink",True,fit=True)
    txt(s,"An unchecked roadmap item is not a release claim. A catalog entry is not universal native qualification.",1.0,6.35,11.3,.38,17,"ink",True,"center")
    footer(s,13); add_notes(s,["This is the credibility slide. It prevents the evolution story from erasing current boundaries.","Sources: docs/user/current-features.md; release-goals.json; PROJECT.md."])

    # 14
    s=slide("ink")
    pill(s,"LITERATE-AI",.76,.55,1.4,"white","ink")
    txt(s,"Readable systems can be\nrigorous, renewable—and fun.",.8,1.55,8.1,1.55,34,"white",True,fit=True)
    txt(s,"Keep intent human-readable.\nKeep implementation replaceable.\nMake evidence earn the release.",.82,3.55,6.4,1.35,21,"fog",fit=True)
    rect(s,.82,5.42,1.1,.07,"orange")
    txt(s,"Start with the visual labs. The grumpy counter has opinions.",.82,5.72,8.8,.36,16,"orange2",True)
    for i,c in enumerate(["orange","blue","green2","orange2","green"]):
        rect(s,9.2+(i%2)*1.45,1.0+(i*1.08),1.2,.45,c,True)
    txt(s,"github.com/jordanhubbard/literate-ai",.82,6.72,5.2,.25,11,"muted",False,font=MONO)
    add_notes(s,["Close on the three imperatives: readable intent, renewable implementation, earned evidence.","Invite attendees to meet the grumpy counter and inspect the generated HTML authority graph."])

    prs.save(OUT)
    digest=hashlib.sha256(OUT.read_bytes()).hexdigest()
    report={"schema":"literate-ai/conference-deck-validation@1","status":"built","slide_count":len(prs.slides),"expected_slide_count":TOTAL,"slide_size_in":[W,H],"speaker_notes_count":sum(bool(sl.notes_slide.notes_text_frame.text.strip()) for sl in prs.slides),"approved_fonts":[FONT,MONO],"sha256":digest,"byte_count":OUT.stat().st_size,"visual_review":"pending render"}
    VALIDATION.write_text(json.dumps(report,indent=2)+"\n",encoding="utf-8")
    print(f"built {OUT} ({len(prs.slides)} slides, {digest})")


if __name__ == "__main__":
    build()
