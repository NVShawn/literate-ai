#!/usr/bin/env python3
"""Verify structural and rendered evidence for the conference deck."""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

from pptx import Presentation

HERE = Path(__file__).resolve().parent
DECK = HERE / "literate-ai-evolution-conference.pptx"
REPORT = HERE / "validation.json"


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--render-dir", type=Path)
    parser.add_argument("--visual-review", choices=("pass", "not-performed"), default="not-performed")
    args = parser.parse_args()

    prs = Presentation(DECK)
    notes = sum(bool(slide.notes_slide.notes_text_frame.text.strip()) for slide in prs.slides)
    fonts: set[str] = set()
    outside: list[dict[str, int]] = []
    for slide_no, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if shape.left < 0 or shape.top < 0 or shape.left + shape.width > prs.slide_width or shape.top + shape.height > prs.slide_height:
                outside.append({"slide": slide_no, "shape_id": shape.shape_id})
            if not shape.has_text_frame:
                continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if run.font.name:
                        fonts.add(run.font.name)

    renders = []
    if args.render_dir:
        renders = sorted(args.render_dir.glob("slide-*.png"))
    checks = {
        "slide_count": len(prs.slides) == 14,
        "speaker_notes": notes == 14,
        "aspect_ratio": abs((prs.slide_width / prs.slide_height) - (16 / 9)) < 0.001,
        "fonts": fonts <= {"Helvetica Neue", "Courier New"},
        "slide_bounds": not outside,
        "render_count": not args.render_dir or len(renders) == 14,
        "visual_review": args.visual_review == "pass",
    }
    digest = hashlib.sha256(DECK.read_bytes()).hexdigest()
    report = {
        "schema": "literate-ai/conference-deck-validation@1",
        "status": "pass" if all(checks.values()) else "fail",
        "checks": checks,
        "slide_count": len(prs.slides),
        "speaker_notes_count": notes,
        "slide_size_emu": [prs.slide_width, prs.slide_height],
        "observed_fonts": sorted(fonts),
        "outside_slide_objects": outside,
        "render_count": len(renders),
        "sha256": digest,
        "byte_count": DECK.stat().st_size,
    }
    REPORT.write_text(json.dumps(report, indent=2) + "\n", encoding="utf-8")
    if report["status"] != "pass":
        raise SystemExit(json.dumps(report, indent=2))
    print(f"verified {len(prs.slides)} slides, {notes} notes, {len(renders)} renders")


if __name__ == "__main__":
    main()
